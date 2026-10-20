"""CV file validation and text extraction (PDF via pdfplumber, DOCX via python-docx)."""
import io
import zipfile

import pdfplumber
from django.conf import settings
from django.core.exceptions import ValidationError
from docx import Document

MIN_TEXT_CHARS = 50


class ScannedDocumentError(ValidationError):
    pass


def detect_kind(name: str, head: bytes, data_len: int, zip_names: list[str] | None = None) -> str:
    """Validate extension + magic bytes and return 'pdf' or 'docx'."""
    lower = name.lower()
    if lower.endswith(".pdf"):
        if not head.startswith(b"%PDF-"):
            raise ValidationError("This file is named .pdf but its content is not a PDF.")
        return "pdf"
    if lower.endswith(".docx"):
        if not head.startswith(b"PK"):
            raise ValidationError("This file is named .docx but its content is not a DOCX document.")
        if zip_names is None or "word/document.xml" not in zip_names:
            raise ValidationError("This file is not a valid DOCX document.")
        return "docx"
    raise ValidationError("Only PDF and DOCX files are accepted.")


def validate_upload(uploaded) -> str:
    """Check size, type and content of an uploaded file. Returns 'pdf' or 'docx'."""
    if uploaded.size > settings.CV_MAX_BYTES:
        raise ValidationError("The file is larger than 5 MB.")
    if uploaded.size == 0:
        raise ValidationError("The file is empty.")
    uploaded.seek(0)
    data = uploaded.read()
    uploaded.seek(0)
    zip_names = None
    if data.startswith(b"PK"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as zf:
                zip_names = zf.namelist()
        except zipfile.BadZipFile:
            raise ValidationError("This file is not a valid DOCX document.")
    return detect_kind(uploaded.name, data[:8], len(data), zip_names)


def extract_pdf_text(fileobj) -> str:
    parts = []
    try:
        with pdfplumber.open(fileobj) as pdf:
            for page in pdf.pages:
                parts.append(page.extract_text() or "")
    except Exception as exc:
        raise ValidationError("Could not read this PDF. It may be corrupt or password-protected.") from exc
    return "\n".join(parts)


def extract_docx_text(fileobj) -> str:
    try:
        doc = Document(fileobj)
    except Exception as exc:
        raise ValidationError("Could not read this DOCX file. It may be corrupt.") from exc
    lines = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                if cell.text not in cells:  # merged cells repeat
                    cells.append(cell.text)
            lines.append(" | ".join(c for c in cells if c.strip()))
    return "\n".join(lines)


def extract_text(fileobj, kind: str) -> str:
    fileobj.seek(0)
    text = extract_pdf_text(fileobj) if kind == "pdf" else extract_docx_text(fileobj)
    text = text.replace("\x00", "").strip()
    if len(text) < MIN_TEXT_CHARS:
        raise ScannedDocumentError(
            "No readable text was found in this file. It looks like a scanned or image-only document. "
            "Please upload a text-based PDF or a DOCX (for example, export it again from Word or Google Docs)."
        )
    return text
